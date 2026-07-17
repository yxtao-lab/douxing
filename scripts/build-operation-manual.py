# -*- coding: utf-8 -*-
"""兜行平台操作手册 .docx 主生成脚本。

按角色分四册（游客 / 发单方 / 商户 / 管理员）+ 附录，引用 docs/操作手册/assets/ 下自动截图。

用法：
  cd project
  python scripts/build-operation-manual.py

前置：
  - 截图已采集到 docs/操作手册/assets/{web,partner,screen,pc,mobile}/
  - 已安装 python-docx（pip install python-docx）
"""
import os
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

# 确保能 import 同目录模块
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from manual_lib import ROOT, OUT  # noqa: E402
import manual_book1
import manual_book2
import manual_book3
import manual_book4
import manual_appendix


def build_cover(doc):
    """构建总封面、目录、前言。"""
    p = doc.add_paragraph()
    p.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(140)
    from docx.shared import RGBColor
    r = p.add_run('兜行平台')
    set_font_local(r, 44, True, (31, 78, 121))
    p2 = doc.add_paragraph()
    p2.alignment = 1
    r2 = p2.add_run('操 作 手 册')
    set_font_local(r2, 40, True, (44, 90, 160))
    p3 = doc.add_paragraph()
    p3.alignment = 1
    p3.paragraph_format.space_before = Pt(30)
    r3 = p3.add_run('（移动端 · PC 用户端 · Web 管理端）')
    set_font_local(r3, 16, False, (90, 90, 90))
    p4 = doc.add_paragraph()
    p4.alignment = 1
    p4.paragraph_format.space_before = Pt(120)
    r4 = p4.add_run('AI 驱动 · 一句话生成专属旅行\n你负责游玩，其他的交给我')
    set_font_local(r4, 14, False, (120, 120, 120))
    p5 = doc.add_paragraph()
    p5.alignment = 1
    p5.paragraph_format.space_before = Pt(80)
    r5 = p5.add_run('版本 v1.0  ·  2026-07-17')
    set_font_local(r5, 12, False, (120, 120, 120))
    doc.add_page_break()

    # 目录
    add_heading_local(doc, '目  录', 0)
    for t in [
        '第一册  游客分册（C 端用户 · 移动端 / PC 用户端）',
        '第二册  发单方分册（C 端发单 · 移动端 / PC 用户端）',
        '第三册  商户分册（B 端 · Web 商户工作台）',
        '第四册  管理员分册（B 端运营 · Web 管理端）',
        '附  录  测试账号、访问地址与常见问题',
    ]:
        pp = doc.add_paragraph()
        pp.paragraph_format.space_after = Pt(4)
        rr = pp.add_run(t)
        set_font_local(rr, 12, True, (0, 0, 0))
    doc.add_page_break()

    # 前言
    add_heading_local(doc, '前  言', 1)
    add_para_local(doc, '兜行（Douxing）定位为 AI 驱动的个性化旅游决策与社交平台，提供 AI 旅行规划、路线、打卡、成就、旅程相册等自助游能力，并扩展发单接单 marketplace（个人/团体发单，旅行社、领队、摄影团队等接单）。平台覆盖三端：')
    add_bullet_local(doc, '移动端：UniApp（H5 / 微信小程序 / Android / iOS），面向 C 端用户。')
    add_bullet_local(doc, 'PC 用户端：Vue3 桌面浏览器，面向 C 端用户与发单方。')
    add_bullet_local(doc, 'Web 管理端：Vue3 + Ant Design Vue，面向 B 端商户与平台运营管理员，含商户工作台 /partner、运营管理、数据中台与运营大屏。')
    add_para_local(doc, '本手册按角色分四册，分别面向游客、发单方、商户、管理员，逐一介绍各功能模块的界面、功能点与操作流程，并配以系统实际截图。')
    pp = doc.add_paragraph()
    pp.paragraph_format.space_before = Pt(4)
    pp.paragraph_format.space_after = Pt(8)
    pPr = pp._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EAF3FB')
    pPr.append(shd)
    rr = pp.add_run('提示  ')
    set_font_local(rr, 10, True, (31, 78, 121))
    rr2 = pp.add_run('本手册截图取自开发环境（zh-CN），实际生产环境界面以最新版本为准；移动端截图以 PC 用户端移动视口示意，功能与移动端完全对齐。')
    set_font_local(rr2, 10, False, (0, 0, 0))
    doc.add_page_break()


# —— 本地简化封装（避免循环 import） ——
def set_font_local(run, size, bold, color):
    from docx.shared import RGBColor
    from docx.oxml.ns import qn
    run.font.name = '微软雅黑'
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '微软雅黑')
    rFonts.set(qn('w:ascii'), '微软雅黑')
    rFonts.set(qn('w:hAnsi'), '微软雅黑')
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_local(doc, text, level):
    sizes = {0: 28, 1: 20, 2: 16, 3: 14, 4: 12}
    colors = {0: (31, 78, 121), 1: (31, 78, 121), 2: (44, 90, 160), 3: (60, 60, 60), 4: (80, 80, 80)}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level <= 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font_local(r, sizes.get(level, 12), True, colors.get(level, (0, 0, 0)))


def add_para_local(doc, text, size=10.5, bold=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.4
    r = p.add_run(text)
    set_font_local(r, size, bold, color)


def add_bullet_local(doc, text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(text)
    set_font_local(r, size, False, None)


def main():
    """生成完整操作手册 .docx。"""
    doc = Document()
    # 默认正文
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    for sec in doc.sections:
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)
        sec.left_margin = Cm(2.4)
        sec.right_margin = Cm(2.4)

    build_cover(doc)
    manual_book1.build(doc)
    manual_book2.build(doc)
    manual_book3.build(doc)
    manual_book4.build(doc)
    manual_appendix.build(doc)

    doc.save(OUT)
    print(f'已生成: {OUT}')
    size = os.path.getsize(OUT) / 1024
    print(f'文件大小: {size:.1f} KB')


if __name__ == '__main__':
    main()
