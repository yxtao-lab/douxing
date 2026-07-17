# -*- coding: utf-8 -*-
"""兜行平台操作手册 .docx 生成工具库。

提供字体、标题、段落、项目符号、编号、截图插入、提示框等通用函数，
以及端子目录常量。各分册 build_*.py 通过 from manual_lib import * 使用。
"""
import os
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
ASSETS = os.path.join(ROOT, 'docs', '操作手册', 'assets')
OUT = os.path.join(ROOT, 'docs', '操作手册', '兜行平台操作手册.docx')

# 端子目录
WEB = 'web'
PARTNER = 'partner'
SCREEN = 'screen'
PC = 'pc'
MOBILE = 'mobile'


def set_font(run, name='微软雅黑', size=None, bold=None, color=None):
    """设置 run 字体（含东亚字体）、字号、加粗、颜色。"""
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    """添加标题（自定义样式，避免 Word 默认蓝色）。"""
    sizes = {0: 28, 1: 20, 2: 16, 3: 14, 4: 12}
    colors = {0: (31, 78, 121), 1: (31, 78, 121), 2: (44, 90, 160), 3: (60, 60, 60), 4: (80, 80, 80)}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level <= 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=sizes.get(level, 12), bold=True, color=colors.get(level, (0, 0, 0)))
    return p


def add_para(doc, text, size=10.5, bold=False, color=None, indent=None):
    """添加正文段落。"""
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.4
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, size=10.5, level=0):
    """添加项目符号段落。"""
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_font(run, size=size)
    return p


def add_number(doc, text, size=10.5):
    """添加编号步骤段落。"""
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_font(run, size=size)
    return p


def add_image(doc, sub, filename, caption, width_cm=14):
    """插入截图 + 图注。sub 为端子目录，filename 为文件名。"""
    path = os.path.join(ASSETS, sub, filename)
    if not os.path.exists(path):
        add_para(doc, f'[截图缺失: {sub}/{filename}]', size=9, color=(200, 0, 0))
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    crun = cap.add_run(caption)
    set_font(crun, size=9, bold=True, color=(90, 90, 90))


def add_note(doc, text):
    """添加提示框（浅色底纹段落）。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Pt(6)
    # 底纹
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EAF3FB')
    pPr.append(shd)
    run = p.add_run('提示  ')
    set_font(run, size=10, bold=True, color=(31, 78, 121))
    run2 = p.add_run(text)
    set_font(run2, size=10)


def page_break(doc):
    """插入分页符。"""
    doc.add_page_break()


def section_title(doc, text):
    """册封面标题。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run(text)
    set_font(run, size=36, bold=True, color=(31, 78, 121))


def section_subtitle(doc, text):
    """册封面副标题。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=16, color=(90, 90, 90))

